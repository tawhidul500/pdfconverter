from __future__ import annotations

import io
import unittest
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont

from app.converter import SourcePdf, analyze_pdf, convert_sources

FIXTURES = Path(__file__).parent / "fixtures"


def digital_pdf() -> bytes:
    return (FIXTURES / "digital.pdf").read_bytes()


def scanned_pdf() -> bytes:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 180), "NovaPDF OCR TEST", fill="black", font=ImageFont.load_default(size=48))
    buffer = io.BytesIO()
    image.save(buffer, format="PDF", resolution=144)
    return buffer.getvalue()


def bengali_scanned_pdf() -> bytes:
    image = Image.open(FIXTURES / "bengali-scan.png").convert("RGB")
    try:
        buffer = io.BytesIO()
        image.save(buffer, format="PDF", resolution=144)
        return buffer.getvalue()
    finally:
        image.close()


class ConverterTests(unittest.TestCase):
    def assert_valid_docx(self, data: bytes) -> Document:
        self.assertTrue(data.startswith(b"PK\x03\x04"))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            self.assertIn("[Content_Types].xml", names)
            self.assertIn("_rels/.rels", names)
            self.assertIn("word/document.xml", names)
            self.assertIsNone(archive.testzip())
        return Document(io.BytesIO(data))

    def test_editable_native_pdf(self) -> None:
        result = convert_sources([SourcePdf("digital.pdf", digital_pdf())], "editable")
        document = self.assert_valid_docx(result.data)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("NovaPDF digital PDF to Word", text)
        self.assertIn("sentence must remain editable", text)
        self.assertEqual(result.page_count, 1)
        self.assertEqual(result.ocr_page_count, 0)

    def test_preserve_layout_pdf(self) -> None:
        result = convert_sources([SourcePdf("digital.pdf", digital_pdf())], "visual")
        document = self.assert_valid_docx(result.data)
        self.assertGreaterEqual(len(document.inline_shapes), 1)


    def test_preserve_layout_with_scanned_ocr_search_layer(self) -> None:
        source = SourcePdf("scan.pdf", scanned_pdf())
        result = convert_sources([source], "visual", use_ocr=True, ocr_language="eng")
        document = self.assert_valid_docx(result.data)
        self.assertGreaterEqual(len(document.inline_shapes), 1)
        text = " ".join(paragraph.text for paragraph in document.paragraphs).upper()
        self.assertIn("NOVAPDF", text)
        self.assertIn("OCR", text)
        self.assertEqual(result.ocr_page_count, 1)
        self.assertEqual(result.native_text_page_count, 0)
        with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("<w:vanish", document_xml)

    def test_preserve_layout_with_native_search_layer(self) -> None:
        result = convert_sources([SourcePdf("digital.pdf", digital_pdf())], "visual", use_ocr=True, ocr_language="eng")
        document = self.assert_valid_docx(result.data)
        self.assertGreaterEqual(len(document.inline_shapes), 1)
        text = " ".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("NovaPDF digital PDF to Word", text)
        self.assertEqual(result.ocr_page_count, 0)
        self.assertEqual(result.native_text_page_count, 1)
        with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("<w:vanish", document_xml)

    def test_preserve_layout_with_bangla_ocr_search_layer(self) -> None:
        source = SourcePdf("bangla-scan.pdf", bengali_scanned_pdf())
        result = convert_sources([source], "visual", use_ocr=True, ocr_language="ben")
        document = self.assert_valid_docx(result.data)
        self.assertGreaterEqual(len(document.inline_shapes), 1)
        text = " ".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("বাংলা", text)
        self.assertIn("পরীক্ষা", text)
        self.assertEqual(result.ocr_page_count, 1)
        with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("<w:vanish", document_xml)

    def test_hybrid_native_layout_uses_editable_positioned_text(self) -> None:
        source = SourcePdf("complex-layout.pdf", (FIXTURES / "complex-layout.pdf").read_bytes())
        result = convert_sources([source], "hybrid", use_ocr=False, ocr_language="eng")
        document = self.assert_valid_docx(result.data)
        self.assertEqual(result.page_count, 1)
        self.assertEqual(result.native_text_page_count, 1)
        self.assertEqual(result.ocr_page_count, 0)
        with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
        self.assertIn("NovaPDFEditableText", document_xml)
        self.assertIn("w:txbxContent", document_xml)
        self.assertIn("Editable Layout Test", document_xml)
        self.assertIn("Alpha", document_xml)
        self.assertNotIn("<w:vanish", document_xml)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(len(document.tables[0].rows), 4)
        self.assertEqual(len(document.tables[0].columns), 3)
        self.assertIn("NovaPDFShape", document_xml)
        self.assertIn("<wp:anchor", document_xml)
        self.assertEqual(len(media), 1)
        self.assertNotIn('behindDoc="1"', document_xml)
        self.assertIn("<w:b", document_xml)
        self.assertIn('w:val="FFFFFF"', document_xml)

    def test_hybrid_line_grid_tables_are_native_and_background_is_clean(self) -> None:
        source = SourcePdf("line-grid-layout.pdf", (FIXTURES / "line-grid-layout.pdf").read_bytes())
        result = convert_sources([source], "hybrid", use_ocr=False, ocr_language="eng")
        document = self.assert_valid_docx(result.data)
        self.assertEqual(result.page_count, 1)
        self.assertEqual(result.native_text_page_count, 1)
        self.assertEqual(result.ocr_page_count, 0)
        self.assertEqual(len(document.tables), 2)
        self.assertEqual((len(document.tables[0].rows), len(document.tables[0].columns)), (4, 2))
        self.assertEqual((len(document.tables[1].rows), len(document.tables[1].columns)), (4, 4))
        self.assertIn("OCR-2026-0730", document.tables[0].cell(0, 1).text)
        self.assertIn("Document scanning service", document.tables[1].cell(1, 0).text)
        with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(media), 1)
            picture = Image.open(io.BytesIO(archive.read(media[0])))
            try:
                self.assertLess(picture.width, 500)
                self.assertLess(picture.height, 300)
            finally:
                picture.close()
        self.assertNotIn('behindDoc="1"', document_xml)
        self.assertIn("NovaPDFShape", document_xml)
        self.assertIn("NovaPDFPicture", document_xml)
        self.assertIn("NovaPDFEditableText", document_xml)
        self.assertIn("<w:b", document_xml)
        self.assertIn('w:val="9C4A00"', document_xml)
        self.assertIn('w:val="36"', document_xml)
        self.assertNotIn("<w:vanish", document_xml)

    def test_hybrid_scanned_layout_uses_visible_ocr_textboxes(self) -> None:
        source = SourcePdf("complex-layout-scan.pdf", (FIXTURES / "complex-layout-scan.pdf").read_bytes())
        result = convert_sources([source], "hybrid", use_ocr=True, ocr_language="eng")
        self.assert_valid_docx(result.data)
        self.assertEqual(result.page_count, 1)
        self.assertEqual(result.native_text_page_count, 0)
        self.assertEqual(result.ocr_page_count, 1)
        with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("NovaPDFEditableText", document_xml)
        self.assertIn("w:txbxContent", document_xml)
        self.assertIn("NovaPDF", document_xml)
        self.assertIn('behindDoc="1"', document_xml)
        self.assertNotIn("<w:vanish", document_xml)

    def test_mixed_orientation_sections(self) -> None:
        source = SourcePdf("mixed.pdf", (FIXTURES / "mixed-orientation.pdf").read_bytes())
        result = convert_sources([source], "visual")
        document = self.assert_valid_docx(result.data)
        self.assertEqual(result.page_count, 2)
        self.assertEqual(len(document.sections), 2)
        first = document.sections[0]
        second = document.sections[1]
        self.assertGreater(first.page_height, first.page_width)
        self.assertGreater(second.page_width, second.page_height)

    def test_scanned_pdf_ocr(self) -> None:
        source = SourcePdf("scan.pdf", scanned_pdf())
        analysis = analyze_pdf(source)
        self.assertEqual(analysis.scanned_pages, (1,))
        result = convert_sources([source], "editable", use_ocr=True, ocr_language="eng")
        document = self.assert_valid_docx(result.data)
        text = " ".join(paragraph.text for paragraph in document.paragraphs).upper()
        self.assertIn("NOVAPDF", text)
        self.assertIn("OCR", text)
        self.assertEqual(result.ocr_page_count, 1)

    def test_scanned_pdf_bangla_ocr(self) -> None:
        source = SourcePdf("bangla-scan.pdf", bengali_scanned_pdf())
        analysis = analyze_pdf(source)
        self.assertEqual(analysis.scanned_pages, (1,))
        result = convert_sources([source], "editable", use_ocr=True, ocr_language="ben")
        document = self.assert_valid_docx(result.data)
        text = " ".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("বাংলা", text)
        self.assertIn("পরীক্ষা", text)
        self.assertEqual(result.ocr_page_count, 1)


if __name__ == "__main__":
    unittest.main()
